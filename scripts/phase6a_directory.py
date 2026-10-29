import csv, json, re, time, html
from pathlib import Path
from datetime import datetime, timezone
from urllib.parse import urlparse
from docx import Document
import requests, feedparser
import argparse

ROOT=Path(__file__).resolve().parent.parent
parser=argparse.ArgumentParser(); parser.add_argument('--probe',action='store_true'); args=parser.parse_args()
doc=Document(ROOT/'docs'/'FrameLine_RSS_Feed_Directory.docx')
paragraphs='\n'.join(p.text for p in doc.paragraphs)
sections={}
for p in doc.paragraphs:
    t=p.text.strip()
    if '  |  ' in t:
        left,right=t.split('  |  ',1); sections[left.strip()]=right.strip()

section_by_table={0:'Global / International',1:'India',2:'United States',3:'United Kingdom & Ireland',4:'France',5:'Germany',6:'Spain',7:'Italy',8:'Japan',9:'China',10:'Russia',11:'Middle East',12:'Africa',13:'Latin America',14:'Canada / Australia / New Zealand',15:'South Korea / Southeast Asia',16:'Science / Technology / Climate'}
rows=[]
for ti,table in enumerate(doc.tables[:17]):
    sec=section_by_table[ti]
    for row in table.rows[1:]:
        publisher, url, note = [html.unescape(c.text.strip()) for c in row.cells]
        lang=(sections.get(sec,'English').split('/')[0].strip().lower())
        # Determine language from directory label and URL/name.
        s=(publisher+' '+sec+' '+url).lower()
        if ti==4: lang='fr'
        elif ti==5: lang='de'
        elif ti==6: lang='es'
        elif ti==7: lang='it'
        elif ti==8 and ('english' in s or 'global newsroom' in s or 'japan times' in s or 'asahi' in s or 'nikkei' in s): lang='en'
        elif ti==8: lang='ja'
        elif ti==9: lang='en' if any(x in s for x in ('english','world (candidate)','cgtn','south china')) else 'zh'
        elif ti==10: lang='en' if any(x in s for x in ('moscow times','rt ')) else 'ru'
        elif ti==11: lang='en' if any(x in s for x in ('national','arab news','daily sabah','anadolu','trt world','iran international')) else 'ar'
        elif ti==12: lang='en' if 'french' not in s else 'fr'
        elif ti==13: lang='en' if 'english' in s else ('pt' if any(x in s for x in ('brasil','folha','g1')) else 'es')
        elif ti==14: lang='en'
        elif ti==15: lang='en' if any(x in s for x in ('korea herald','cna','straits times')) else 'ko'
        if 'VOA' in publisher:
            note += '; dropped: obviously corrupted/incomplete URL (encoded endpoint path); not probed'
        rows.append({'publisher':publisher,'section':sec,'url':url,'language':lang,'region':sec,'directory_status_note':note})

with open(ROOT/'docs'/'candidate_feeds.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=['publisher','section','url','language','region','directory_status_note']);w.writeheader();w.writerows(rows)

cfg=json.loads((ROOT/'data'/'rss_feeds.json').read_text(encoding='utf-8'))
existing={x['url'] for group in cfg.values() for x in group}
want={
'Deutsche Welle':['https://rss.dw.com/rdf/rss-en-all'],'France 24':['https://www.france24.com/en/rss'],'RFI':['https://www.rfi.fr/en/rss'],'Euronews':['https://www.euronews.com/rss'],'NHK World':['https://www3.nhk.or.jp/rssxml/news/globalnewsroom.xml'],'The Japan Times':['https://www.japantimes.co.jp/feed/'],'Asahi Shimbun':['https://www.asahi.com/ajw/rss/'],'Daily Sabah':['https://www.dailysabah.com/rssFeed/world'],'Arab News':['https://www.arabnews.com/rss.xml'],'The National':['https://www.thenationalnews.com/arc/outboundfeeds/rss/'],'Anadolu Agency':['https://www.aa.com.tr/en/rss/default?cat=world'],'TRT World':['https://www.trtworld.com/rss'],'Daily Maverick':['https://www.dailymaverick.co.za/dmrss/'],'Premium Times':['https://www.premiumtimesng.com/feed'],'Africanews':['https://www.africanews.com/feed/rss'],'ABC Australia':['https://www.abc.net.au/news/feed/45910/rss.xml'],'CBC':['https://www.cbc.ca/webfeed/rss/rss-world'],'SBS News':['https://www.sbs.com.au/news/feed'],'RNZ':['https://www.rnz.co.nz/rss/news.xml'],'Korea Herald':['https://www.koreaherald.com/rss/'],'CNA':['https://www.channelnewsasia.com/api/v1/rss-outbound-feed?_format=xml'],'RTÉ':['https://www.rte.ie/feeds/rss/?index=/news/'],'Agência Brasil':['https://agenciabrasil.ebc.com.br/en/rss/ultimasnoticias/feed.xml'],
'tagesschau':['https://www.tagesschau.de/index~rss2.xml'],'ZEIT Online':['https://newsfeed.zeit.de/all'],'Le Monde':['https://www.lemonde.fr/international/rss_full.xml'],'franceinfo':['https://www.francetvinfo.fr/monde.rss'],'El País':['https://feeds.elpais.com/mrss-s/pages/ep/site/elpais.com/section/internacional/portada'],'ANSA':['https://www.ansa.it/sito/notizie/mondo/mondo_rss.xml'],'Folha de S.Paulo':['https://feeds.folha.uol.com.br/emcimadahora/rss091.xml'],'El Universal':['https://www.eluniversal.com.mx/rss.xml'],'Al Arabiya':['https://www.alarabiya.net/tools/rss'],
'Hindi publisher (The Hindu Hindi)':[]}
# Hindi feed candidate absent from DOCX: The Hindu Hindi not available; use Amar Ujala public RSS lead.
want['Amar Ujala']=['https://www.amarujala.com/rss/india-news.xml']
selected=[]
for name,urls in want.items():
    for url in urls:
        if url not in existing:
            selected.append((name,url))

print('CANDIDATES',len(rows),'NEW_TO_PROBE',len(selected))
for name,url in selected: print(name+'\t'+url)

if args.probe:
    headers={'User-Agent':'FrameLine/1.0 (academic news-framing research prototype; https://github.com/Sanidhya-06/FrameLine)'}
    report=[]
    for name,url in selected:
        print('ROUND1',name,flush=True)
        try:
            t=time.monotonic(); r=requests.get(url,headers=headers,timeout=15); elapsed=round(time.monotonic()-t,2)
            entries=feedparser.parse(r.content).entries if r.status_code==200 else []
            report.append({'publisher':name,'url':url,'status1':r.status_code,'entries1':entries,'time1':elapsed,'raw1':r.content if r.status_code==200 else b''})
        except Exception as e: report.append({'publisher':name,'url':url,'status1':type(e).__name__,'entries1':[],'time1':15,'raw1':b''})
    print('WAITING 60 seconds between fetch rounds',flush=True); time.sleep(60)
    for item in report:
        print('ROUND2',item['publisher'],flush=True)
        try:
            t=time.monotonic(); r=requests.get(item['url'],headers=headers,timeout=15); elapsed=round(time.monotonic()-t,2)
            entries=feedparser.parse(r.content).entries if r.status_code==200 else []
            item.update(status2=r.status_code,entries2=entries,time2=elapsed,raw2=r.content if r.status_code==200 else b'')
        except Exception as e: item.update(status2=type(e).__name__,entries2=[],time2=15,raw2=b'')
    def shares(entries):
        n=len(entries)
        if not n:return (0,0,0,0)
        return tuple(round(100*sum(bool(e.get(k)) for e in entries)/n,2) for k in ('title','link','summary','published_parsed'))
    compact=[]
    for x in report:
        es=x.get('entries2',[]); n=len(es); ages=[]
        for e in es:
            try:
                from datetime import datetime,timezone
                stamp=datetime(*e.published_parsed[:6],tzinfo=timezone.utc).timestamp()
                ages.append((datetime.now(timezone.utc).timestamp()-stamp)/3600)
            except Exception: pass
        compact.append({k:x.get(k) for k in ('publisher','url','status1','status2','time1','time2')}|{'entries1':len(x['entries1']),'entries2':n,'title_pct':shares(es)[0],'url_pct':shares(es)[1],'summary_pct':shares(es)[2],'date_pct':shares(es)[3],'newest_age_h':round(min(ages),2) if ages else ''})
        if x.get('raw2') and n>=10 and shares(es)[0]>=90 and shares(es)[1]>=90 and x.get('status1')==200 and x.get('status2')==200:
            clean=urlparse(x['url']).netloc.replace('.','_')+'_'+urlparse(x['url']).path.strip('/').replace('/','_')
            clean=re.sub(r'[^A-Za-z0-9_-]','_',clean)[:70]
            (ROOT/'tests'/'fixtures'/f'feed_phase6a_{clean}.xml').write_bytes(x['raw2'])
    rawout=[]
    for x in report:
        rawout.append({**{k:v for k,v in x.items() if k not in ('entries1','entries2','raw1','raw2')},'entries1':[{'title':e.get('title',''),'link':e.get('link',''),'summary':e.get('summary',''),'published':e.get('published',''),'published_parsed':list(e.get('published_parsed') or [])} for e in x['entries1']],'entries2':[{'title':e.get('title',''),'link':e.get('link',''),'summary':e.get('summary',''),'published':e.get('published',''),'published_parsed':list(e.get('published_parsed') or [])} for e in x.get('entries2',[])],'raw2':x.get('raw2',b'').decode('utf-8','replace')})
    (ROOT/'docs'/'phase6a_probe_raw.json').write_text(json.dumps(rawout,ensure_ascii=False),encoding='utf-8')
    with open(ROOT/'docs'/'phase6a_probe.csv','w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(compact[0]));w.writeheader();w.writerows(compact)
